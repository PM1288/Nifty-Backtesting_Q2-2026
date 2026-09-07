#!/usr/bin/env python3
"""Validate supplied source data and print an idempotent transactional SQL import.

No credentials, network login, schema truncation or publication-time fabrication.
Run with the existing extraction Python environment (pandas, xlrd, openpyxl).
SQL output is an artifact; pipe to psql with ON_ERROR_STOP only after validation.
"""
import argparse
import hashlib
import importlib.util
import json
import sys
import shutil
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("parsers", ROOT / "services/nse_fii_reports_service/src/nse_fii_services/parsers.py")
parsers = importlib.util.module_from_spec(spec)
spec.loader.exec_module(parsers)

def sql(v):
    if v is None:
        return "NULL"
    return "'" + str(v).replace("'", "''") + "'"

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source-dir", type=Path, required=True)
    ap.add_argument("--validate-only", action="store_true")
    ap.add_argument("--archive-dir", type=Path, help="Required for SQL generation; content-addressed immutable source copies")
    args = ap.parse_args()
    p = args.source_dir
    golden = json.loads((p / "Trading_Analytics_Golden_Reference_20260904_v1_0.json").read_text())
    xls = p / "fii_stats_04-Sep-2026.xls"
    workbook = p / "FII_OI_ Stats_July_2023onwards.xlsx"
    frame = parsers.parse_fii_stats_excel(xls.read_bytes())
    assert len(frame) == 16, f"Expected 16 product rows, got {len(frame)}"
    columns = ["buy_contracts", "buy_value_in_Cr", "sell_contracts", "sell_value_in_Cr", "open_contracts", "open_contracts_value_in_Cr"]
    expected_keys = ["buy_contracts", "buy_crore", "sell_contracts", "sell_crore", "oi_contracts", "oi_crore"]
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    stats_cells = list(wb["FII_Stats_Data"].iter_rows(min_row=8149, max_row=8167, max_col=7, values_only=True))
    for expected in golden["fii_derivatives_records"]:
        r = frame[frame.fii_derivatives == expected["product"]].iloc[0]
        cells = stats_cells[expected["workbook_row"] - 8149]
        assert cells[0].strip() == expected["product"]
        for i, (col, key) in enumerate(zip(columns, expected_keys)):
            assert Decimal(str(r[col])) == Decimal(str(expected[key])) == Decimal(str(cells[i+1])), (expected["product"], col)
    # The raw CSV was not supplied. Import only the verified final workbook block,
    # explicitly labelled WORKBOOK_VERIFIED_BLOCK, not as an observed official CSV.
    sheet = wb["FII_OI_Data"]
    cells = list(sheet.iter_rows(min_row=4119, max_row=4125, max_col=15, values_only=True))
    assert "Sep 04" in str(cells[0][0])
    participant_columns = [str(v).strip().lower().replace(" ", "_") for v in cells[1][1:]]
    rows = cells[2:]
    for row, expected in zip(rows, golden["participant_oi_records"]):
        assert row[0] == expected["participant"]
        assert dict(zip(participant_columns, row[1:])) == expected["values"]
    wb.close()
    print(json.dumps({"stats_rows":16,"xls_and_workbook_fields_matched":96,"participant_workbook_fields_matched":70,"participant_source":"WORKBOOK_VERIFIED_BLOCK_NOT_RAW_CSV","unknown_year_blocks":"QUARANTINED_NOT_IMPORTED","known_at":"actual import time, never report date"}), file=sys.stderr)
    if args.validate_only:
        return
    if not args.archive_dir:
        ap.error("--archive-dir is required when generating an import")
    args.archive_dir.mkdir(parents=True, exist_ok=True)
    now = datetime.now(timezone.utc).isoformat()
    report_date = golden["report_date"]
    statements = ["BEGIN;", "SELECT pg_advisory_xact_lock(2026090701);"]
    for path, kind, count in [(xls, "DATED_FII_XLS", 16), (workbook, "WORKBOOK_VERIFIED_BLOCK", 5)]:
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        archived = args.archive_dir / (sha + path.suffix)
        if archived.exists():
            assert hashlib.sha256(archived.read_bytes()).hexdigest() == sha, "Archive checksum conflict"
        else:
            shutil.copyfile(path, archived)
        path = archived
        identity = "trading-analytics:" + kind + ":" + sha
        metadata = json.dumps({"publication_time":"UNKNOWN", "source_rows":"A4:G22" if count==16 else "FII_OI_Data!A4121:O4125", "date_resolution":"DATED_FILENAME" if count==16 else "FINAL_BLOCK_MATCHES_SUPPLIED_GOLDEN_20260904", "raw_csv_supplied":False})
        statements.append("INSERT INTO audit.trading_analytics_artifacts (artifact_id,report_date,source_kind,source_path,sha256,retrieved_at,published_at,known_at,parser_version,row_count,metadata) VALUES (" + ",".join(sql(v) for v in [identity,report_date,kind,str(path),sha,now,None,now,"TRADING-ANALYTICS-20260907.0",count,metadata]) + ") ON CONFLICT DO NOTHING;")
        if count == 16:
            table="market_data.nse_fii_derivatives_stats"
            keys=["fii_derivatives"]+[c.lower() for c in columns]
            records=[[r.fii_derivatives]+[str(int(r[c])) if j%2==0 else format(Decimal(str(r[c])),".2f") for j,c in enumerate(columns)] for _,r in frame.iterrows()]
        else:
            table="market_data.nse_fii_participant_open_interest"
            keys=["client_type"]+participant_columns
            records=[list(row) for row in rows]
        for record in records:
            cols=["run_id","run_kind","loaded_at","trade_date","source_file"]+keys
            values=[identity,"verified_source_import",now,report_date,str(path)]+record
            statements.append(f"INSERT INTO {table} ({','.join(cols)}) SELECT {','.join(sql(v) for v in values)} WHERE NOT EXISTS (SELECT 1 FROM {table} WHERE run_id={sql(identity)} AND {keys[0]}={sql(record[0])});")
    statements.append("COMMIT;")
    print("\n".join(statements))

if __name__ == "__main__":
    main()
